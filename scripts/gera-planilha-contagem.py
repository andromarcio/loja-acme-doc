#!/usr/bin/env python3
"""gera-planilha-contagem.py — planilha de entrega da contagem (Excel).

Gera um .xlsx genérico com um processo elementar por linha, no layout que as
equipes de métrica usam na entrega: Requisito · Processos elementares · Tipo ·
Qtd. INM · TD/DER (Qtd., Descrição) · RLR/ALR (Qtd., Descrição) — mais uma
coluna final `Item do Jira`, com a chave do ticket que alcançou a feature.

A coluna do Jira vai no **fim** de propósito: o bloco anterior é o layout que a
equipe de métrica já usa, e deslocar as colunas dela quebraria a conferência.
A chave sai das AIMs (`analise-impacto/AIM-*.md`) — a mesma fonte que a auditoria por
sprint usa: na AIM do ticket, o `ticket:` do front-matter É a chave e a
`## Alterações na spec, por Feature Set` lista as features alcançadas; na AIM da sprint,
a chave vem da coluna `Ticket`. Feature que nenhum ticket alcançou sai com `—`; feature
alcançada por mais de um sai com as chaves separadas por espaço.

`Requisito` é o **Feature Set (N2)** a que a feature pertence — o nome, sem o
código, em TitleCase e com extensão `.docx` (ex.: `AtendimentoSocial.docx`),
lido do `# Feature Set: X` do README ao lado do N3.

A fonte é sempre o N3 — `## Métricas de tamanho` para os números e, para as
descrições de DER e ALR, o bloco ```json de cada processo elementar na
`### Memória de cálculo` (`{"pe", "alr", "der", "nao_contados", "motivo"}`). Só o
bloco: a enumeração em prosa misturava campo e comentário, e a planilha copiava os
dois (memória antiga → `migra-enumeracao.py`). Nada é recalculado aqui: a
planilha ESPELHA a contagem gravada na feature (ver a regra "A contagem nasce no
N3" no CLAUDE.md). Sem o bloco, as colunas Descrição saem vazias — o número
existe, a enumeração não.

A linha de **0 PF** fica na planilha, com o motivo À VISTA na Descrição do DER
(`Não contado: …`, da memória), e não num comentário de célula (decisão do PO,
2026-09-28): o PO precisa ver que a feature foi impactada e não foi contada. Sem o
motivo na memória, a célula diz que ele falta. Linha com `—` no PF (ainda não
medida, ou PE reutilizado `↪`) não entra.

`Qtd. INM` sai em branco por definição: é preenchida pela equipe de métrica.

Geração **sob demanda**: rode quando for entregar a contagem. A planilha é saída,
não fonte — não se versiona (está no .gitignore) e não é gerada em CI; o que roda
em CI é o teste deste script sobre `scripts/__fixtures__/planilha-contagem/`.
Mudou a contagem? Corrija o N3 e gere de novo.

Uso:
  python3 scripts/gera-planilha-contagem.py [raiz-da-instância] [-o saida.xlsx]
                                            [--escopo PREFIXO ...] [--jira [CHAVE ...]]
                                            [--der-alr]

  --escopo  limita as TRANSAÇÕES às features cujo ID começa pelo prefixo (ex.:
            TEM-ASO, CEP); a aba de funções de dados sai inteira
  --jira    com chaves (ex.: STRY0012345, ISSUE-482, PDTIC25093-49): só o que elas
            alcançaram, nas duas abas; sem valor: a sprint inteira — tudo o que alguma
            AIM cita, inclusive o que foi entregue sem ticket
  --der-alr recorte de CONFERÊNCIA, não formato de entrega: esconde Tipo e
            Qtd. INM para isolar a enumeração de DER e de ALR. A planilha que
            vai para a equipe de métricas é a padrão, sem esta opção

A coluna **Natureza** — última de cada aba, nos dois layouts — diz se o processo elementar (ou o arquivo lógico) foi
**Incluído** ou **Alterado** pela sprint — o que separa o que conta 100% em PFB
do que conta 50% em PFL. Sai das AIMs, pela coluna `Natureza`; a AIM da sprint é a
única fonte do que foi entregue **sem ticket**, e por isso entra na leitura junto com
as AIMs dos tickets.
"""
import argparse, json, re, sys, unicodedata
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("✗ openpyxl não encontrado. Instale com: pip install openpyxl")

# ── leitura de tabelas markdown, sempre guiada pelo CABEÇALHO ────────────────
# Posição de coluna varia entre instâncias; nome de coluna, não. Ler por posição
# devolve o número errado com cara de número certo.
def celulas(linha):
    return [c.strip() for c in linha.strip().strip("|").split("|")]

def eh_separador(linha):
    return re.fullmatch(r"\|[\s\-:|]+\|", linha.strip()) is not None

def tabelas(linhas, ini, parar):
    """Todas as tabelas a partir de `ini` até `parar`; cada uma como (cab, linhas)."""
    out, cab, corpo = [], None, []
    for l in linhas[ini:]:
        if parar(l):
            break
        if not l.strip().startswith("|"):
            if cab is not None:
                out.append((cab, corpo)); cab, corpo = None, []
            continue
        if eh_separador(l):
            continue
        if cab is None:
            cab = celulas(l)
        else:
            corpo.append(celulas(l))
    if cab is not None:
        out.append((cab, corpo))
    return out

def col(cab, *nomes):
    for i, c in enumerate(cab):
        for n in nomes:
            if re.fullmatch(n, c, re.I):
                return i
    return -1

def secao(linhas, titulo):
    for i, l in enumerate(linhas):
        if re.fullmatch(rf"#{{2}}\s+{titulo}\s*", l.strip()):
            return i
    return -1

def corpo_da_secao(linhas, titulo):
    """Linhas de `## titulo` até o próximo `## ` — os `###` dela vêm junto."""
    i = secao(linhas, titulo)
    if i < 0:
        return []
    fim = next((j for j in range(i + 1, len(linhas)) if re.match(r"##\s", linhas[j])), len(linhas))
    return linhas[i + 1:fim]

def chave_pe(s):
    """Nome do PE para casar a tabela com a memória: sem crase, negrito nem caixa."""
    return re.sub(r"\s+", " ", re.sub(r"[`*]", "", s)).strip(" .;").casefold()

# ── extração por feature ────────────────────────────────────────────────────
def id_da_feature(raw):
    l = next((x for x in raw.splitlines() if re.search(r"N[íi]vel 3", x)), "")
    m = re.search(r"`([A-Z]{3}-[A-Z]{3}-\d{2})`", l)
    return m.group(1) if m else None

def title_case(s):
    """"Atendimento Social" -> "AtendimentoSocial"; acentos e pontuação saem."""
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    palavras = [w for w in re.split(r"[^A-Za-z0-9]+", s) if w]
    return "".join(w[0].upper() + w[1:] for w in palavras)

def nome_do_n2(caminho):
    """Nome do Feature Set (N2), do `# Feature Set: X` do README irmão do N3."""
    readme = caminho.parent / "README.md"
    if readme.exists():
        for l in readme.read_text(encoding="utf-8").splitlines()[:12]:
            m = re.match(r"#\s+Feature Set:\s*(.+?)\s*$", l)
            if m:
                return m.group(1)
            if l.startswith("# "):  # instância que não usa o prefixo
                return re.sub(r"^#\s+", "", l).strip()
    return None

def requisito(caminho):
    """Coluna Requisito: o N2 em TitleCase com extensão .docx, sem o código."""
    n2 = nome_do_n2(caminho)
    if not n2:
        return None
    return f"{title_case(n2)}.docx"

def enumeracao(sec):
    """Os blocos ```json da `### Memória de cálculo` → `{chave do PE: bloco}`.

    `sec` é o corpo de `## Métricas de tamanho`. A enumeração de ALR e DER é DADO, não
    prosa — `{"pe": …, "alr": […], "der": […]}`, mais `nao_contados` (vira comentário
    da célula) e, na linha de 0 PF, `motivo` —, e é só ela que a planilha lê. A memória
    em prosa dizia campo e comentário na mesma frase, e não havia como separar um do
    outro: a explicação do CPM saía como o primeiro ALR, o parêntese de anotação como
    parte do nome do DER (planilha do portal-compras, 2026-10-02). O formato é cobrado
    ao gravar o N3 (`valida-enumeracao-contagem.mjs`); a memória antiga se converte com
    `migra-enumeracao.py`. Bloco com JSON inválido é como bloco ausente.
    """
    blocos, corpo, na_memoria, dentro = {}, [], False, False
    for l in sec:
        s = l.strip()
        if re.match(r"###\s+Mem[óo]ria de c[áa]lculo", s):
            na_memoria = True
        elif na_memoria and not dentro and s == "```json":
            dentro, corpo = True, []
        elif dentro and s == "```":
            dentro = False
            try:
                b = json.loads("\n".join(corpo))
            except ValueError:
                continue
            if isinstance(b, dict) and isinstance(b.get("pe"), str):
                blocos[chave_pe(b["pe"])] = b
        elif dentro:
            corpo.append(l)
    return blocos

def processos(caminho):
    raw = caminho.read_text(encoding="utf-8")
    fid = id_da_feature(raw)
    if not fid:
        return []
    linhas = raw.splitlines()
    i = secao(linhas, "Métricas de tamanho")
    if i < 0:
        return []
    tabs = tabelas(linhas, i + 1, lambda l: re.match(r"#{2,3}\s", l) is not None)
    if not tabs:
        return []
    cab, corpo = tabs[0]
    iPE = col(cab, r"Fun[çc][ãa]o de Transa[çc][ãa]o", r"Processo elementar")
    iTipo, iALR, iDER, iPF = col(cab, r"Tipo"), col(cab, r"ALR"), col(cab, r"DER"), col(cab, r"PF")
    if iPE < 0 or iPF < 0:
        return []
    blocos = enumeracao(corpo_da_secao(linhas, "Métricas de tamanho"))
    req = requisito(caminho)
    linhas_out = []
    for c in corpo:
        if iPF >= len(c) or not re.fullmatch(r"\d+", c[iPF]):
            continue  # `—` = ainda não medida
        pega = lambda k: c[k] if 0 <= k < len(c) else ""
        nome = re.sub(r"[`*]", "", pega(iPE)).strip()
        b = blocos.get(chave_pe(nome))
        lista = lambda k: [str(x) for x in ((b or {}).get(k) or [])]
        alr_d, der_d = lista("alr"), lista("der")
        nota = f"Não contados: {b['nao_contados']}" if b and b.get("nao_contados") else ""
        # PF 0: a linha fica, com o motivo À VISTA na Descrição do DER, e não num
        # comentário que só aparece ao passar o mouse (decisão do PO, 2026-09-28): o
        # PO precisa ver que a feature foi impactada e NÃO foi contada — uma linha só
        # de traços parece item esquecido. Sem o motivo, a linha diz que ele falta.
        if pega(iPF) == "0":
            motivo = re.sub(r"\s+", " ", str((b or {}).get("motivo") or "")).strip()
            der_d = [f"Não contado: {motivo}" if motivo else "Não contado — a memória de cálculo não diz por quê"]
            alr_d, nota = [], ""
        linhas_out.append({
            "requisito": req, "pe": pega(iPE), "tipo": pega(iTipo),
            "der_qtd": pega(iDER), "der_desc": "\n".join(der_d),
            "alr_qtd": pega(iALR), "alr_desc": "\n".join(alr_d),
            "nota_alt": nota,
            "tem_memoria": b is not None,
            "id": fid,
        })
    return linhas_out

# ── rastreabilidade: feature → chave do ticket ──────────────────────────────
# A chave NÃO vive no N3 (lá mora a `## Origem`, que aponta para a AIM do ticket). Quem
# liga feature a ticket, com a natureza (incluída/alterada), é a AIM (decisão do PO,
# 2026-09-28): a do ticket — `analise-impacto/AIM-<CHAVE>.md` — lista na
# `## Alterações na spec, por Feature Set` as features que ele alcançou; a da sprint
# consolida a entrega, inclusive o que foi entregue sem ticket. O `tipo` do
# front-matter separa as duas, e o `ticket:` é a chave — sem adivinhar pelo nome.
PASTA_AIM = "analise-impacto"
SECAO_ALTERACOES = "Alterações na spec, por Feature Set"
SECAO_DADOS = "Funções de dados alteradas"
# Chave de ticket: a mesma lista do STORY_KEY_RE de scripts/lib/trace-index.mjs
# (STRY…, ISSUE-n, PDTIC…-n, EXP-…) mais o formato Jira (PRJ25001-49). Na coluna
# Ticket da AIM da sprint, a chave em crase vale em qualquer formato.
STORY_KEY = r"(?:STRY\d{4,}|ISSUE-\d{1,7}|PDTIC\d+-\d{1,7}|EXP-[A-Za-z0-9][A-Za-z0-9_-]*)"
JIRA_LEGADO = r"(?:[A-Za-z]+\d+-\d+)"
CHAVE_TICKET = rf"(?:{STORY_KEY}|{JIRA_LEGADO})"

def front_matter(texto):
    """{campo: valor} do front-matter (carimbo e linhas em branco antes do `---` valem;
    o comentário de fim de linha não faz parte do valor)."""
    linhas, i, out = texto.splitlines(), 0, {}
    while i < len(linhas) and (not linhas[i].strip() or re.match(r"^\s*<!--.*-->\s*$", linhas[i])):
        i += 1
    if i >= len(linhas) or linhas[i].strip() != "---":
        return out
    for l in linhas[i + 1:]:
        if l.strip() == "---":
            break
        m = re.match(r"^([a-z][\w-]*):\s*(.*)$", l, re.I)
        if m:
            out[m.group(1).lower()] = re.sub(r"\s+#.*$", "", m.group(2)).strip().strip("\"'")
    return out

def aims(raiz):
    """[(chave | None, Path)] — chave preenchida = AIM do ticket; None = AIM da sprint."""
    out = []
    for arq in sorted((raiz / PASTA_AIM).glob("AIM-*.md")):
        fm = front_matter(arq.read_text(encoding="utf-8"))
        tipo = fm.get("tipo", "").lower()
        if tipo == "ticket" and fm.get("ticket") and not fm["ticket"].startswith("["):
            out.append((fm["ticket"], arq))
        elif tipo == "sprint":
            out.append((None, arq))
    return out

def secao_md(texto, titulo):
    """O corpo de `## <titulo>` até o próximo `## ` — "" se a seção não existe."""
    m = re.search(rf"^##\s+{re.escape(titulo)}\s*$", texto, re.M)
    if not m:
        return ""
    prox = re.search(r"^##\s", texto[m.end():], re.M)
    return texto[m.end(): m.end() + prox.start()] if prox else texto[m.end():]

def mapa_tickets(raiz):
    """{fid: {"jira": [chaves], "natureza": "incluída"|"alterada"|""}}

    Duas fontes, porque nenhuma sozinha cobre a sprint:

      • a AIM do ticket — a chave é o `ticket:` do front-matter e a
        `## Alterações na spec, por Feature Set` lista as features alcançadas;
      • a AIM da sprint — a única fonte das features **entregues sem ticket**:
        ler só as AIMs dos tickets deixa a planilha com features a menos, com a
        mesma cara de estar completa.

    Nos dois casos a leitura é pelo CABEÇALHO. A coluna `Natureza` diz se o
    processo elementar foi **incluído** (conta 100% em PFB) ou **alterado**
    (conta 50% em PFL) — é o que separa as duas metades de um projeto de
    melhoria, e sem ela a métrica não fecha o cálculo.
    """
    mapa = {}

    def registra(fid, chave, natureza):
        d = mapa.setdefault(fid, {"jira": [], "natureza": ""})
        if chave and chave not in d["jira"]:
            d["jira"].append(chave)
        if natureza:
            d["natureza"] = natureza

    def limpa_nat(x):
        v = re.sub(r"[^a-zà-ú]", "", re.sub(r"[`*]", "", x).strip().lower())
        return {"incluida": "incluída", "nova": "incluída"}.get(v, v) \
            if v in ("alterada", "incluída", "incluida", "nova") else ""

    for chave, arq in aims(raiz):
        if chave is None:
            continue
        trecho = secao_md(arq.read_text(encoding="utf-8"), SECAO_ALTERACOES)
        nat_por_id = {}
        for cab, corpo in tabelas(trecho.splitlines(), 0, lambda l: False):
            iF, iN = col(cab, r"^Feature$"), col(cab, r"^Natureza$")
            if iF < 0 or iN < 0:
                continue
            for c in corpo:
                m = re.match(r"`([A-Z]{3}-[A-Z]{3}-\d{2})`", c[iF]) if iF < len(c) else None
                if m and iN < len(c):
                    nat_por_id[m.group(1)] = limpa_nat(c[iN])
        for fid in re.findall(r"^\| `([A-Z]{3}-[A-Z]{3}-\d{2})` \*\*", trecho, re.M):
            registra(fid, chave, nat_por_id.get(fid, ""))

    # A AIM da sprint fecha a entrega: traz as features sem ticket e confirma a
    # natureza das demais. A chave sai da coluna `Ticket` — quando ela não é uma
    # chave ("⚠️ sem ticket", "—"), a feature entra sem chave, que é a informação
    # verdadeira, e não fica de fora da contagem.
    for chave, arq in aims(raiz):
        if chave is not None:
            continue
        linhas = secao_md(arq.read_text(encoding="utf-8"), SECAO_ALTERACOES).splitlines()
        for cab, corpo in tabelas(linhas, 0, lambda l: False):
            iF, iN = col(cab, r"^Feature$"), col(cab, r"^Natureza$")
            iJ = col(cab, r"^Ticket$", r"Item do Jira")
            if iF < 0 or iN < 0:
                continue
            for c in corpo:
                m = re.match(r"`([A-Z]{3}-[A-Z]{3}-\d{2})`", c[iF]) if iF < len(c) else None
                if not m:
                    continue
                # A coluna É a da chave: vale o que estiver em crase, em qualquer
                # formato, e todas — célula com duas chaves traz as duas. Sem crase,
                # só o que tem cara de chave; um ID de feature citado não é chave.
                chaves = []
                if 0 <= iJ < len(c):
                    chaves = [k for k in (re.findall(r"`([^`\s]+)`", c[iJ])
                                          or re.findall(rf"(?<![\w-]){CHAVE_TICKET}(?![\w-])", c[iJ], re.I))
                              if not re.fullmatch(r"[A-Z]{3}-[A-Z]{3}-\d{2}", k)]
                for chave in chaves or [""]:
                    registra(m.group(1), chave, limpa_nat(c[iN]) if iN < len(c) else "")
    return mapa

# `secao()` casa `##` com o título inteiro — é o que o caminho das transações
# precisa. Os cabeçalhos do DATA-MODEL são `###` e trazem texto depois do termo
# ("### ALIs — Arquivos Lógicos Internos"), então este caminho usa um casador
# próprio. Ampliar o compartilhado consertaria aqui e mudaria o comportamento lá.
def secoes_amplas(linhas, *termos):
    """Índices de TODOS os cabeçalhos que citam algum dos termos.

    Devolve todos, e não o primeiro, porque o documento costuma ter um cabeçalho
    de seção sem tabela ("## Arquivos Lógicos (APF)") seguido do que tem a tabela
    ("### ALIs — …"). Parar no primeiro devolve zero linha com cara de "não há
    funções de dados".
    """
    return [i for i, l in enumerate(linhas)
            if re.match(r"#{2,3}\s", l.strip())
            and any(re.search(t, l.strip(), re.I) for t in termos)]

# ── funções de dados (ALI / AIE) ────────────────────────────────────────────
# A contagem de dados não vive no N3 — vive no DATA-MODEL (índice, com RLR/DER/PF)
# e nos fragmentos por domínio (com as entidades constituintes, que são o que a
# memória oferece como descrição de RLR). Sem esta aba a planilha entrega só
# metade da contagem: as transações.
NOME_ARQ_LOGICO = (r"^ALI$", r"^AIE$", r"ALI\s*/\s*AIE", r"Arquivo l[óo]gico")

def tipo_da_tabela(cab, titulo):
    """ALI ou AIE pela coluna do nome e, na falta, pelo título da seção. Numa tabela
    mista ("ALI / AIE | Tipo | …") devolve "": lá o tipo vem da coluna Tipo, linha a linha."""
    if col(cab, r"^AIE$") >= 0:
        return "AIE"
    if col(cab, r"^ALI$") >= 0:
        return "ALI"
    ali = re.search(r"\bALIs?\b|Internos?\b", titulo, re.I)
    aie = re.search(r"\bAIEs?\b|Interface Externa", titulo, re.I)
    return "AIE" if aie and not ali else "ALI" if ali and not aie else ""

def registro_ali_aie(raiz):
    """{nome: (tipo, domínio)} do `global/ALI-AIE-MAP.md`, o registro canônico."""
    arq = raiz / "global" / "ALI-AIE-MAP.md"
    if not arq.is_file():
        return {}
    lim = lambda x: re.sub(r"\s+", " ", re.sub(r"[`*]", "", x)).strip(" .;")
    reg = {}
    for cab, corpo in tabelas(arq.read_text(encoding="utf-8").splitlines(), 0, lambda l: False):
        iN, iT, iD = col(cab, r"^Nome$"), col(cab, r"^Tipo$"), col(cab, r"Dom[íi]nio")
        if iN < 0:
            continue
        for c in corpo:
            nome = lim(c[iN]) if iN < len(c) else ""
            if not nome or nome.startswith("["):
                continue
            tipo = lim(c[iT]).upper() if 0 <= iT < len(c) else ""
            dom = lim(c[iD]) if 0 <= iD < len(c) else ""
            reg[nome] = (tipo if tipo in ("ALI", "AIE") else "", "" if dom in ("—", "-") else dom)
    return reg

def funcoes_de_dados(raiz):
    idx = raiz / "global" / "DATA-MODEL.md"
    if not idx.is_file():
        return []
    linhas = idx.read_text(encoding="utf-8").splitlines()
    # O índice do template traz DUAS tabelas — "### ALIs — …" (ALI | Domínio | …) e
    # "### AIEs — …" (AIE | Sistema externo | Entidades / estruturas usadas | …) — e há
    # instância com uma tabela só, mista, com coluna Tipo. Lê todas: parar na primeira
    # entregava a aba sem as AIEs, com a mesma cara de "não há AIE".
    achadas = []
    for i in secoes_amplas(linhas, r"\bAL[IE]s?\b", r"Arquivos L[óo]gicos", r"Interface Externa"):
        for cab, corpo in tabelas(linhas, i + 1, lambda l: re.match(r"#{2,3}\s", l) is not None):
            if col(cab, *NOME_ARQ_LOGICO) >= 0 and col(cab, r"RLR") >= 0:
                achadas.append((cab, corpo, tipo_da_tabela(cab, linhas[i])))
    if not achadas:
        return []

    # entidades constituintes e CAMPOS: o fragmento por domínio é quem os nomeia.
    # Cada entidade abre com "> **ALI: X** · papel" (ou "> **AIE: X** · estrutura
    # externa de …"), que é o elo entidade→arquivo lógico; a tabela logo abaixo lista
    # os campos em Label PO. O título do fragmento dá o domínio de quem o declara.
    const, campos, dom_frag = {}, {}, {}
    for frag in sorted((raiz / "global" / "data-models").glob("*.md")):
        flin = frag.read_text(encoding="utf-8").splitlines()
        mdom = next((m for m in (re.match(r"#\s+Data Model:\s*(.+?)\s*$", l) for l in flin[:5]) if m), None)
        for k, l in enumerate(flin):
            m = re.match(r"^>\s*\*\*(?:ALI|AIE):\s*([^*]+)\*\*", l)
            if not m:
                continue
            if mdom:
                dom_frag.setdefault(m.group(1).strip(), mdom.group(1))
            ali = m.group(1).strip()
            j = k + 1
            while j < len(flin) and not flin[j].startswith("| Label PO"):
                if re.match(r"^##\s", flin[j]):
                    break
                j += 1
            if j >= len(flin) or not flin[j].startswith("| Label PO"):
                continue
            vistos = campos.setdefault(ali, [])
            for linha in flin[j + 2:]:
                if not linha.startswith("|"):
                    break
                campo = re.sub(r"[`*]", "", celulas(linha)[0]).strip()
                if campo and campo not in vistos:
                    vistos.append(campo)

        for j in secoes_amplas(flin, r"Arquivos L[óo]gicos deste dom[íi]nio"):
            for fcab, fcorpo in tabelas(flin, j + 1, lambda l: re.match(r"#{2,3}\s", l) is not None):
                k, kc = col(fcab, r"ALI\s*/\s*AIE"), col(fcab, r"Entidades constituintes")
                if k >= 0 and kc >= 0:
                    for c in fcorpo:
                        if k < len(c) and kc < len(c):
                            nome = re.sub(r"[`*]", "", c[k]).strip()
                            const[nome] = c[kc]
                            if mdom:
                                dom_frag.setdefault(nome, mdom.group(1))
                    break

    reg = registro_ali_aie(raiz)
    lim = lambda x: re.sub(r"\s+", " ", re.sub(r"[`*]", "", x)).strip(" .;")
    saida, vistos = [], set()
    for cab, corpo, tipo_tab in achadas:
        iNome = col(cab, *NOME_ARQ_LOGICO)
        iDom, iRLR, iDER, iPF = col(cab, r"Dom[íi]nio"), col(cab, r"RLR"), col(cab, r"DER"), col(cab, r"PF")
        iTipo = col(cab, r"^Tipo$")
        iEnt = col(cab, r"Entidades constituintes", r"Entidades\s*/\s*estruturas.*")
        for c in corpo:
            nome = lim(c[iNome]) if iNome < len(c) else ""
            if not nome or nome.lower().startswith("total") or nome in vistos:
                continue
            vistos.add(nome)
            pega = lambda k: lim(c[k]) if 0 <= k < len(c) else ""
            reg_tipo, reg_dom = reg.get(nome, ("", ""))
            # O tipo decide o peso (ALI 7/10/15, AIE 5/7/10): vem da coluna Tipo, da
            # tabela ou do registro, nesta ordem. Sem nenhum dos três, ALI — o que o
            # script sempre assumiu.
            tipo = pega(iTipo).upper()
            if tipo not in ("ALI", "AIE"):
                tipo = tipo_tab or reg_tipo or "ALI"
            # A tabela de AIEs não tem coluna Domínio (tem "Sistema externo"): sem
            # ela, o domínio vem do registro e, na falta, do fragmento que declara o
            # arquivo lógico.
            dom = pega(iDom) if iDom >= 0 else (reg_dom or dom_frag.get(nome, ""))
            # A descrição de RLR são as entidades constituintes, UMA POR LINHA. O
            # data-model as escreve em dois níveis — "Principal (principal) · A, B, C
            # (subgrupos)" —, então não basta quebrar no `·`: os subgrupos ficariam
            # todos numa linha só e a coluna deixaria de ser enumeração. Quebra-se
            # também na vírgula. Cuidado: nomes como "Critério/Decisão de Desempate +
            # Inscrição da Decisão" trazem `+` e `/` e NÃO se quebram — só a vírgula
            # separa itens neste formato. Sem o fragmento, vale a coluna do índice.
            ents = const.get(nome, "") or pega(iEnt)
            ents = [lim(e) for grupo in re.split(r"\s+·\s+",
                        re.sub(r"\((principal|subgrupos?|suporte)\)", "", ents))
                    for e in grupo.split(",") if lim(e)]
            saida.append({
                "requisito": dom, "pe": nome, "tipo": tipo,
                "der_qtd": pega(iDER), "der_desc": "\n".join(campos.get(nome, [])),
                "alr_qtd": pega(iRLR), "alr_desc": "\n".join(ents),
                "tipo_registro": reg_tipo,
                "id": nome,
            })
    return saida

# quais arquivos lógicos cada ticket alterou — `## Funções de dados alteradas` das
# AIMs. Traz também o RLR/DER DEPOIS da alteração: o CPM dimensiona a função alterada
# (CHGA) pelo tamanho que ela tem DEPOIS, e o DATA-MODEL guarda o antes. Entregar
# só o antes dá à métrica o número errado para o cálculo que ela vai fazer.
def natureza_dados(raiz):
    """{arquivo lógico: "incluída"|"alterada"} — da tabela de totais da AIM da sprint.

    A coluna `Natureza` das tabelas por ALI fala das TABELAS ("Tabela incluída"),
    não da função de dados: um ALI alterado pode ganhar uma tabela nova sem
    deixar de ser alterado. A natureza da FUNÇÃO só existe na tabela de totais.
    """
    out = {}
    for chave, arq in aims(raiz):
        if chave is not None:
            continue
        linhas = arq.read_text(encoding="utf-8").splitlines()
        for cab, corpo in tabelas(linhas, 0, lambda l: False):
            iN = col(cab, r"Natureza da .*fun[çc][ãa]o.*")
            iF = col(cab, r"Fun[çc][ãa]o de dados")
            if iF < 0 or iN < 0:
                continue
            for c in corpo:
                if iF >= len(c) or iN >= len(c):
                    continue
                nome = re.sub(r"[`*]", "", c[iF]).strip()
                v = re.sub(r"[^a-zà-ú]", "", re.sub(r"[`*]", "", c[iN]).strip().lower())
                if nome and not nome.lower().startswith("total") and v in ("alterada", "incluída", "incluida"):
                    out[nome] = "incluída" if v.startswith("inclu") else "alterada"
    return out

def mapa_tickets_dados(raiz):
    mapa, depois = {}, {}
    for chave, arq in aims(raiz):
        if chave is None:
            continue
        trecho = secao_md(arq.read_text(encoding="utf-8"), SECAO_DADOS)
        for cabec in re.findall(r"^###\s+(?:ALI|AIE):\s*(.+)$", trecho, re.M):
            nome = cabec.split("—")[0].strip()
            if not nome or nome.startswith("["):
                continue
            mapa.setdefault(nome, [])
            if chave not in mapa[nome]:
                mapa[nome].append(chave)
            mrlr = re.search(r"RLR\s+(\d+)\s*→\s*(\d+)", cabec)
            mder = re.search(r"DER\s+(\d+)\s*→\s*(\d+)", cabec)
            if mrlr or mder:
                depois[nome] = {
                    "rlr": (mrlr.group(1), mrlr.group(2)) if mrlr else None,
                    "der": (mder.group(1), mder.group(2)) if mder else None,
                }
    return mapa, depois

# ── planilha ────────────────────────────────────────────────────────────────
CINZA = PatternFill("solid", fgColor="D9D9D9")
FINA = Side(style="thin", color="000000")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
LARGURAS = [42.2, 44.7, 12.2, 6.2, 8.0, 27.2, 5.8, 34.5, 18.0, 12.0]

# A coluna fala do processo elementar (e do arquivo lógico), não da feature de
# onde a natureza veio — por isso "Incluído/Alterado", no masculino.
NATUREZA = {"incluída": "Incluído", "alterada": "Alterado"}

def cabecalho(ws, rot_a, rot_b, rot_ef, rot_gh, rot_f="Descrição", rot_h="Descrição"):
    ws.merge_cells("A1:A2"); ws["A1"] = rot_a
    ws.merge_cells("B1:B2"); ws["B1"] = rot_b
    ws.merge_cells("C1:C2"); ws["C1"] = "Tipo"
    ws.merge_cells("D1:D2"); ws["D1"] = "Qtd.\nINM"
    ws.merge_cells("E1:F1"); ws["E1"] = rot_ef
    ws.merge_cells("G1:H1"); ws["G1"] = rot_gh
    ws.merge_cells("I1:I2"); ws["I1"] = "Item do Jira"
    ws.merge_cells("J1:J2"); ws["J1"] = "Natureza"
    ws["E2"], ws["F2"], ws["G2"], ws["H2"] = "Qtd.", rot_f, "Qtd.", rot_h
    for r in (1, 2):
        for c in range(1, 11):
            cel = ws.cell(r, c)
            cel.font = Font(name="Arial", size=10, bold=True)
            cel.fill = CINZA
            cel.border = BORDA
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(LARGURAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A3"

# Layout enxuto — só a enumeração de DER e de ALR. Serve à conferência que a
# equipe de métricas faz campo a campo: Tipo, Qtd. INM e complexidade são ruído
# quando a pergunta é "quais elementos vocês reconheceram aqui". A identificação
# (requisito/domínio, processo/arquivo lógico e a chave do Jira) fica — sem ela a
# lista de campos não se prende a nada.
LARGURAS_DER_ALR = [40.0, 44.0, 7.0, 46.0, 7.0, 40.0, 18.0, 12.0]

def cabecalho_der_alr(ws, rot_a, rot_b, rot_cd, rot_ef, rot_d, rot_f):
    ws.merge_cells("A1:A2"); ws["A1"] = rot_a
    ws.merge_cells("B1:B2"); ws["B1"] = rot_b
    ws.merge_cells("C1:D1"); ws["C1"] = rot_cd
    ws.merge_cells("E1:F1"); ws["E1"] = rot_ef
    ws.merge_cells("G1:G2"); ws["G1"] = "Item do Jira"
    ws.merge_cells("H1:H2"); ws["H1"] = "Natureza"
    ws["C2"], ws["D2"], ws["E2"], ws["F2"] = "Qtd.", rot_d, "Qtd.", rot_f
    for r in (1, 2):
        for c in range(1, 9):
            cel = ws.cell(r, c)
            cel.font = Font(name="Arial", size=10, bold=True)
            cel.fill = CINZA
            cel.border = BORDA
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(LARGURAS_DER_ALR, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A3"

def corpo_der_alr(ws, linhas):
    r = 3
    for d in linhas:
        for c, v in enumerate([d["requisito"], d["pe"],
                               d["der_qtd"], d["der_desc"],
                               d["alr_qtd"], d["alr_desc"],
                               d.get("jira") or "—",
                               NATUREZA.get(d.get("natureza"), "—")], start=1):
            cel = ws.cell(r, c)
            if v is not None and v != "":
                cel.value = int(v) if isinstance(v, str) and v.isdigit() else v
            cel.font = Font(name="Arial", size=10)
            cel.border = BORDA
            cel.alignment = Alignment(
                horizontal="center" if c in (3, 5, 7, 8) else "left",
                vertical="top", wrap_text=True)
            if c == 3 and d.get("nota_alt"):
                cel.comment = Comment(d["nota_alt"], "docqui", height=90, width=320)
        r += 1
    return r - 3

def corpo(ws, linhas):
    r = 3
    for d in linhas:
        for c, v in enumerate([d["requisito"], d["pe"], d["tipo"], None,
                               d["der_qtd"], d["der_desc"], d["alr_qtd"], d["alr_desc"],
                               d.get("jira") or "—",
                               NATUREZA.get(d.get("natureza"), "—")], start=1):
            cel = ws.cell(r, c)
            if v is not None and v != "":
                cel.value = int(v) if isinstance(v, str) and v.isdigit() else v
            cel.font = Font(name="Arial", size=10)
            cel.border = BORDA
            cel.alignment = Alignment(
                horizontal="center" if c in (3, 4, 5, 7, 9, 10) else "left",
                vertical="top", wrap_text=True)
            if c == 5 and d.get("nota_alt"):
                cel.comment = Comment(d["nota_alt"], "docqui", height=90, width=320)
        r += 1
    return r - 3

def escreve(linhas, saida, dados=None, so_der_alr=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Funções de Transação"

    if so_der_alr:
        cabecalho_der_alr(ws, "Requisito", "Processos elementares", "TD / DER", "RLR / ALR",
                          "Elementos de dado (campo a campo)", "Arquivos lógicos referenciados")
        n = corpo_der_alr(ws, linhas)
    else:
        cabecalho(ws, "Requisito", "Processos elementares", "TD / DER", "RLR / ALR")
        n = corpo(ws, linhas)

    # A contagem tem duas metades. Entregar só as transações é entregar metade —
    # e quem confere o total não encontra os 117 PF das funções de dados.
    if dados:
        wd = wb.create_sheet("Funções de Dados")
        # O rótulo distingue as duas populações: a Qtd. é o DER que o baseline
        # contou; a lista é o inventário de campos do data-model. Chamar as duas
        # de "Descrição" convida a ler a segunda como enumeração da primeira.
        if so_der_alr:
            cabecalho_der_alr(wd, "Domínio", "Arquivo lógico (ALI / AIE)", "TD / DER", "RLR / RET",
                              "Campos (data-model), campo a campo", "Entidades constituintes")
            corpo_der_alr(wd, dados)
        else:
            cabecalho(wd, "Domínio", "Arquivo lógico (ALI / AIE)", "TD / DER", "RLR / RET",
                      rot_f="Campos (data-model)", rot_h="Entidades constituintes")
            corpo(wd, dados)

    wb.save(saida)
    return n

def main():
    ap = argparse.ArgumentParser(description="Planilha de entrega da contagem (Excel), espelhada dos N3.")
    ap.add_argument("raiz", nargs="?", default=".", help="raiz da instância (padrão: a pasta atual)")
    ap.add_argument("-o", "--saida", default=None,
                    help="arquivo .xlsx (padrão: contagem-processos-elementares.xlsx na raiz)")
    ap.add_argument("--escopo", nargs="*", default=None, metavar="PREFIXO",
                    help="limita as transações às features com este prefixo de ID; "
                         "a aba de funções de dados sai inteira")
    ap.add_argument("--jira", nargs="*", default=None, metavar="CHAVE",
                    help="só o que estes tickets alcançaram nas AIMs (STRY…, ISSUE-n, "
                         "PDTIC…-n, EXP-…, Jira), nas duas abas; sem valor, a sprint "
                         "inteira — inclusive o que foi entregue sem ticket")
    ap.add_argument("--der-alr", action="store_true",
                    help="recorte de conferência: esconde Tipo e Qtd. INM para isolar a\n"
                         "enumeração de DER e de ALR. NÃO é o formato de entrega")
    a = ap.parse_args()

    raiz = Path(a.raiz).resolve()
    if not (raiz / "modules").is_dir():
        sys.exit(f"✗ Não parece uma instância (sem modules/): {raiz}")
    saida = Path(a.saida) if a.saida else raiz / "contagem-processos-elementares.xlsx"

    linhas = []
    for p in sorted(raiz.glob("modules/**/f-*.md")):
        linhas += processos(p)
    impacto = mapa_tickets(raiz)
    for d in linhas:
        info = impacto.get(d["id"])
        d["jira"] = " ".join(info["jira"]) if info else ""
        # A natureza é da FEATURE; o PE herda. Um PE nasce dentro de uma feature
        # incluída ou dentro de uma alterada — não há terceiro caso no delta.
        d["natureza"] = (info or {}).get("natureza", "")
        d["na_sprint"] = info is not None
    contadas = len(linhas)
    if a.escopo:
        linhas = [d for d in linhas if any(d["id"].startswith(e.upper()) for e in a.escopo)]
    no_escopo = len(linhas)
    if a.jira is not None:
        alvo = {k.upper() for k in a.jira}
        # Sem chaves informadas: a sprint inteira — inclusive o que foi entregue
        # sem item no Jira. Filtrar por `d["jira"]` deixaria essas de fora.
        linhas = [d for d in linhas
                  if (d["na_sprint"] and (not alvo or any(k.upper() in alvo for k in d["jira"].split())))]
    linhas.sort(key=lambda d: (d["id"], d["pe"]))

    # Planilha vazia pode ter três causas, e cada uma se corrige num lugar diferente:
    # dizer só "nenhum processo contado" mandava procurar o erro no N3 quando o que
    # esvaziou a lista foi o filtro — ou uma AIM que o script não reconheceu.
    if not linhas:
        if not contadas:
            sys.exit("✗ Nenhum processo elementar contado encontrado (coluna PF preenchida no N3).")
        if not no_escopo:
            sys.exit(f"✗ {contadas} processo(s) contado(s), nenhum com o prefixo {' '.join(a.escopo)}.")
        rels = aims(raiz)
        chaves = sorted({k for info in impacto.values() for k in info["jira"]})
        sys.exit(f"✗ {no_escopo} processo(s) contado(s), nenhum alcançado "
                 f"{'por ' + ' '.join(a.jira) if a.jira else 'por AIM'}. "
                 f"AIMs lidas ({PASTA_AIM}/): {sum(1 for k, _ in rels if k)} "
                 f"de ticket, {sum(1 for k, _ in rels if not k)} de sprint; "
                 f"chaves encontradas: {', '.join(chaves) or 'nenhuma'}.")

    dados = funcoes_de_dados(raiz)
    jira_d, depois = mapa_tickets_dados(raiz)
    nat_d = natureza_dados(raiz)
    for d in dados:
        d["jira"] = " ".join(jira_d.get(d["pe"], []))
        d["natureza"] = nat_d.get(d["pe"], "")
        # Na sprint = alguma AIM a cita: a do ticket (Funções de dados alteradas) ou a
        # da sprint, que é a única fonte da função de dados alterada sem ticket — o
        # mesmo critério das transações.
        d["na_sprint"] = bool(d["jira"] or d["natureza"])
        alt = depois.get(d["pe"])
        if alt:
            # a célula passa a levar o tamanho DEPOIS; o antes fica na descrição
            nota = []
            for chv, campo in (("rlr", "alr_qtd"), ("der", "der_qtd")):
                if alt[chv]:
                    antes, dps = alt[chv]
                    d[campo] = dps
                    nota.append(f"{chv.upper()} {antes} → {dps}")
            # A nota NÃO entra na enumeração: aquela coluna é a lista de DER e
            # nada mais — recado no meio dos campos faz o primeiro item parecer um
            # deles. Vira anotação da célula de Qtd., que é o número a que ela se
            # refere: quem confere o tamanho passa o mouse e lê; quem lê a lista
            # de campos não tropeça.
            if nota:
                d["nota_alt"] = ("Alterado por " + (d["jira"] or "item sem chave") + ": " + " · ".join(nota)
                                 + ". A quantidade é a DEPOIS da alteração, base do CHGA.")
    if a.jira is not None:
        alvo = {k.upper() for k in a.jira}
        dados = [d for d in dados
                 if (d["na_sprint"] and (not alvo or any(k.upper() in alvo for k in d["jira"].split())))]

    n = escreve(linhas, saida, dados, so_der_alr=a.der_alr)
    print(f"✓ {n} processo(s) elementar(es) e {len(dados)} função(ões) de dados em {saida}")
    if not dados:
        print("⚠️  Nenhuma função de dados — a contagem sai pela metade. "
              "A fonte é global/DATA-MODEL.md (tabelas de ALIs e de AIEs) + global/data-models/*.md.")
    # É o tipo que escolhe a tabela de peso (ALI 7/10/15, AIE 5/7/10). A planilha fica
    # com o do índice — foi com ele que o PF do índice foi calculado —, mas índice e
    # registro discordando é erro de um dos dois, e não cabe à planilha arbitrar.
    tipo_dif = [d for d in dados if d.get("tipo_registro") and d["tipo_registro"] != d["tipo"]]
    if tipo_dif:
        print(f"⚠️  {len(tipo_dif)} arquivo(s) lógico(s) com tipo diferente no DATA-MODEL e no "
              f"ALI-AIE-MAP — o tipo escolhe o peso; corrija a fonte errada:")
        for d in tipo_dif[:8]:
            print(f"      {d['pe']}: {d['tipo']} no DATA-MODEL · {d['tipo_registro']} no ALI-AIE-MAP")
    # A enumeração de campos NÃO é a lista de DER que o baseline contou: o
    # data-model lista atributos físicos da engenharia reversa, e o DER do CPM é
    # elemento reconhecido pelo usuário, sem repetição, considerando todos os PE.
    # Divergir é o esperado — o que não pode é passar despercebido.
    dif = []
    for d in dados:
        if str(d["der_qtd"]).isdigit():
            n = len([x for x in d["der_desc"].split("\n") if x.strip() and not x.startswith("[")])
            if n != int(d["der_qtd"]):
                dif.append((d["pe"], d["der_qtd"], n))
    if dif:
        print(f"ℹ️  {len(dif)} arquivo(s) lógico(s) em que a quantidade de DER do baseline difere "
              f"do número de campos do data-model — esperado, são populações diferentes:")
        for pe, q, n in dif:
            print(f"      {pe}: DER {q} no baseline · {n} campos no data-model")

    # O RLR do baseline e as entidades constituintes descrevem a mesma coisa por
    # dois caminhos: quantos subgrupos reconhecíveis o arquivo lógico tem, e quais
    # são eles. Divergir não prova erro — o CPM permite agrupar várias entidades
    # físicas num único RET —, mas a lista deixa de sustentar o número, e é isso
    # que a auditoria vai perguntar. Sinalizar, não arbitrar.
    rlr_dif = []
    for d in dados:
        if str(d["alr_qtd"]).isdigit() and d["alr_desc"]:
            n_ent = len([x for x in d["alr_desc"].split("\n") if x.strip()])
            if n_ent != int(d["alr_qtd"]):
                rlr_dif.append((d["pe"], d["alr_qtd"], n_ent))
    if rlr_dif:
        print(f"⚠️  {len(rlr_dif)} arquivo(s) lógico(s) em que o RLR não bate com as entidades "
              f"constituintes enumeradas — a lista não sustenta o número. Pode ser agrupamento "
              f"legítimo de entidades num mesmo RET; confirme com a equipe de métricas:")
        for pe, q, n_ent in rlr_dif:
            print(f"      {pe}: RLR {q} · {n_ent} entidade(s) enumerada(s)")

    sem_ent = [d for d in dados if not d["alr_desc"]]
    if sem_ent:
        print(f"⚠️  {len(sem_ent)} função(ões) de dados sem entidades constituintes — coluna "
              f"Descrição de RLR vazia: {', '.join(d['pe'] for d in sem_ent[:6])}")

    # Natureza em branco vira "—" na célula, que é indistinguível de "não se
    # aplica". Num recorte de sprint toda linha tem de ter uma: é ela que decide
    # se o processo entra como PFB (incluído) ou PFL (alterado).
    sem_nat = [d for d in linhas + dados if d.get("na_sprint", True) and not d.get("natureza")]
    if sem_nat:
        print(f"⚠️  {len(sem_nat)} linha(s) sem natureza (incluído/alterado) — a coluna sai `—`, "
              f"que a métrica não consegue distinguir de 'não se aplica'. A fonte é a coluna "
              f"`Natureza` das AIMs:")
        for d in sem_nat[:8]:
            print(f"      {d.get('id', d['pe'])} — {d['pe']}")

    sem_n2 = [d for d in linhas if not d["requisito"]]
    if sem_n2:
        print(f"⚠️  {len(sem_n2)} sem N2 resolvido — coluna Requisito vazia. "
              f"O nome sai do `# Feature Set: X` do README ao lado do N3.")
        for d in sem_n2[:6]:
            print(f"      {d['id']} — {d['pe']}")

    sem_mem = [d for d in linhas if not d["tem_memoria"]]
    if sem_mem:
        print(f"⚠️  {len(sem_mem)} sem a enumeração na memória de cálculo — colunas Descrição vazias. "
              f"Ela é um bloco ```json por processo elementar na ### Memória de cálculo do N3; "
              f"memória antiga: python3 scripts/migra-enumeracao.py.")
        for d in sem_mem[:6]:
            print(f"      {d['id']} — {d['pe']}")

    # Memória que não bate com o número não sustenta o número: se a quantidade
    # declarada difere do que a descrição enumera, uma das duas está errada.
    def itens(s):
        return len([x for x in s.split("\n") if x.strip()])
    divergem = []
    for d in linhas:
        if not d["der_desc"] and not d["alr_desc"]:
            continue
        for rot, qtd, desc in (("DER", d["der_qtd"], d["der_desc"]), ("ALR", d["alr_qtd"], d["alr_desc"])):
            if qtd.isdigit() and itens(desc) != int(qtd):
                divergem.append(f"{d['id']} — {d['pe']}: {rot} diz {qtd}, a memória enumera {itens(desc)}")
    if divergem:
        print(f"⚠️  {len(divergem)} divergência(s) entre a quantidade e o que a memória enumera:")
        for x in divergem[:8]:
            print(f"      {x}")
        print("      Corrija no N3 — a planilha espelha a fonte, não a conserta.")

if __name__ == "__main__":
    main()
