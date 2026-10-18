#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera um documento Word (.docx) com a especificação de um Feature Set (N2):
- Visão geral do N2 (a partir do README do Feature Set)
- Para cada Feature (N3), o conteúdo negocial: tudo, menos o front-matter, os
  comentários HTML, os blocos dev-only e o Changelog; a seção "Métricas de tamanho"
  sai enquanto nenhuma linha foi contada
- Os protótipos existentes, embutidos como imagem
- Na capa, o sistema da instância (Sigla e Nome do global/MASTER.md)

Fonte de verdade dos Markdowns: modules/<dominio>/<feature-set>/.
As imagens (fluxo do N2 e protótipos) são pré-renderizadas — ver render_*.mjs.

Uso:
  python3 gen_feature_set_docx.py \
      --fs-dir modules/parametrizacao/maquina-estado \
      --img-dir exports/_assets/maquina-estado \
      --out exports/PAR-MAE__Maquina-de-Estado.docx

Convenção de imagens no --img-dir:
  overview.png            -> diagrama do Fluxo Principal do N2 (substitui o mermaid)
  proto-<ID>.png          -> protótipo da feature <ID> (ex.: proto-PAR-MAE-07.png)
  proto-<ID>-<n>.png      -> protótipo adicional da mesma feature (ordem por <n>)
"""
import argparse
import os
import re
import glob
import datetime
import tempfile

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

# ---------------------------------------------------------------- layout
PAGE_TEXT_WIDTH_IN = 6.9      # Letter - 0.8in margens
PAGE_AVAIL_HEIGHT_IN = 8.9    # altura útil por página para uma imagem
CODE_SHADE = "F4F5F7"
HEAD_SHADE = "E8EEF4"
MONO_FONT = "Consolas"
DEV_ONLY_MARKER = '<div class="dev-only">'

# ================================================================ parsing

def strip_inline_tags(text):
    text = re.sub(r"</?small>", "", text)
    text = re.sub(r"</?strong>", "", text)
    text = re.sub(r"</?em>", "", text)
    return text


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator_row(cells):
    return all(re.match(r"^:?-{2,}:?$", c.strip()) for c in cells if c.strip() != "") and any(
        c.strip() for c in cells
    )


def parse_blocks(lines):
    """Converte um trecho de Markdown numa lista de blocos tipados."""
    blocks = []
    i, n = 0, len(lines)
    while i < n:
        raw = lines[i]
        s = raw.strip()
        if s == "":
            i += 1
            continue
        if s.startswith("<!--"):
            i += 1
            continue
        # code fence
        if s.startswith("```"):
            lang = s[3:].strip()
            body = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            blocks.append(("code", (lang, body)))
            continue
        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("heading", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue
        # horizontal rule
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            i += 1
            continue
        # table
        if s.startswith("|"):
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            rows = [split_table_row(r) for r in tbl]
            rows = [r for r in rows if not is_separator_row(r)]
            if rows:
                blocks.append(("table", rows))
            continue
        # blockquote
        if s.startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]).rstrip("\n"))
                i += 1
            blocks.append(("quote", quote))
            continue
        # list
        if re.match(r"^\s*([-*]|\d+\.)\s+", raw):
            items = []
            while i < n and re.match(r"^\s*([-*]|\d+\.)\s+", lines[i]):
                lm = re.match(r"^\s*([-*]|\d+\.)\s+(.*)$", lines[i])
                ordered = bool(re.match(r"^\s*\d+\.", lines[i]))
                items.append((ordered, lm.group(2).strip()))
                i += 1
            blocks.append(("list", items))
            continue
        # paragraph
        para = []
        while i < n:
            l = lines[i]
            ls = l.strip()
            if (
                ls == ""
                or ls.startswith("|")
                or ls.startswith(">")
                or ls.startswith("#")
                or ls.startswith("```")
                or re.match(r"^(-{3,}|\*{3,}|_{3,})$", ls)
                or re.match(r"^\s*([-*]|\d+\.)\s+", l)
            ):
                break
            para.append(ls)
            i += 1
        blocks.append(("para", " ".join(para)))
    return blocks


def business_lines(path):
    """Linhas negociais do N3: sem o front-matter, os comentários HTML e os blocos
    dev-only. O bloco dev-only pode estar em qualquer ponto do N3 — no N3 anterior à
    6.4.16, a `## Superfície` tinha o seu, logo depois da Origem —, então ele é pulado
    até o `</div>` que o fecha, e a leitura continua depois dele."""
    with open(path, encoding="utf-8") as fh:
        text = fh.read()
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = text.split("\n")
    inicio = next((i for i, l in enumerate(lines) if l.strip()), None)
    if inicio is not None and lines[inicio].strip() == "---":
        fim = next((i for i in range(inicio + 1, len(lines)) if lines[i].strip() == "---"), None)
        if fim is not None:
            lines = lines[fim + 1:]
    out, depth = [], 0
    for line in lines:
        if depth == 0 and DEV_ONLY_MARKER not in line:
            out.append(line.rstrip("\r"))
            continue
        depth += line.count("<div") - line.count("</div>")
    return out


def secoes_h2(lines):
    """Quebra as linhas em [(título ou None, linhas)] pela rubrica `## ` no início da
    linha, fora de cerca de código (a rubrica entra nas linhas da própria seção)."""
    secoes, cerca = [(None, [])], False
    for l in lines:
        if l.strip().startswith(("```", "~~~")):
            cerca = not cerca
        m = None if cerca else re.match(r"^##\s+(.+?)\s*$", l)
        if m:
            secoes.append((m.group(1).strip(), [l]))
        else:
            secoes[-1][1].append(l)
    return secoes


COLUNAS_CONTADAS = ("Tipo", "ALR", "DER", "Complexidade", "PF")
CELULA_VAZIA = {"", "—", "–", "-"}


def metricas_vazia(linhas):
    """True quando a seção "Métricas de tamanho" não tem nada contado: nenhuma célula
    de Tipo, ALR, DER, Complexidade ou PF preenchida e nenhuma Memória de cálculo. A
    linha de PF 0 (PE descartado) e a de PE reutilizado (↪ no Tipo) contam como
    preenchidas. Tabela sem nenhuma dessas colunas não é julgada vazia: na dúvida, a
    seção fica."""
    if any(re.match(r"^###\s+Mem[óo]ria de c[áa]lculo", l.strip()) for l in linhas):
        return False
    cab = None
    for l in linhas:
        if not l.strip().startswith("|"):
            cab = None
            continue
        cells = [strip_inline_tags(c) for c in split_table_row(l)]
        if is_separator_row(cells):
            continue
        if cab is None:
            cab = cells
            if not any(c in cab for c in COLUNAS_CONTADAS):
                return False
            continue
        for nome in COLUNAS_CONTADAS:
            if nome in cab:
                k = cab.index(nome)
                if k < len(cells) and cells[k] not in CELULA_VAZIA:
                    return False
    return True


def secoes_exportadas(lines):
    """Tira do N3 o Changelog (o histórico fica no repositório) e a seção "Métricas de
    tamanho" enquanto nada foi contado."""
    out = []
    for titulo, linhas in secoes_h2(lines):
        if titulo and re.match(r"^Changelog\b", titulo, re.I):
            continue
        if titulo and re.match(r"^M[ée]tricas de tamanho$", titulo, re.I) and metricas_vazia(linhas):
            continue
        out.extend(linhas)
    return out

# ================================================================ docx render helpers

def set_cell_shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]*\))")


def add_inline(paragraph, text):
    """Adiciona runs com **negrito**, `mono` e [texto](link)->texto."""
    text = strip_inline_tags(text)
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(9.5)
        elif tok.startswith("["):
            lm = re.match(r"\[([^\]]+)\]\([^)]*\)", tok)
            paragraph.add_run(lm.group(1))
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_cell_content(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    segments = text.split("<br>")
    for idx, seg in enumerate(segments):
        if idx > 0:
            p.add_run().add_break()
        add_inline(p, seg.strip())


def render_table(doc, rows):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx in range(ncols):
            val = row[cidx] if cidx < len(row) else ""
            add_cell_content(cells[cidx], val)
            if ridx == 0:
                set_cell_shade(cells[cidx], HEAD_SHADE)
                for par in cells[cidx].paragraphs:
                    for r in par.runs:
                        r.bold = True
        if ridx == 0:
            set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_code(doc, lang, body):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.rows[0].cells[0]
    set_cell_shade(cell, CODE_SHADE)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    lines = body if body else [""]
    for idx, ln in enumerate(lines):
        if idx > 0:
            p.add_run().add_break()
        r = p.add_run(ln if ln != "" else " ")
        r.font.name = MONO_FONT
        r.font.size = Pt(8.5)
        # comentários gherkin em cinza
        if ln.lstrip().startswith("#"):
            r.font.color.rgb = RGBColor(0x6A, 0x73, 0x7D)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_quote(doc, quote_lines):
    text = " ".join(l for l in quote_lines if l.strip() != "")
    if not text.strip():
        return
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    add_inline(p, text)
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(9.5)
        if r.font.color.rgb is None:
            r.font.color.rgb = RGBColor(0x55, 0x5F, 0x66)


def render_list(doc, items):
    for ordered, text in items:
        style = "List Number" if ordered else "List Bullet"
        p = doc.add_paragraph(style=style)
        add_inline(p, text)


def add_image_fit(doc, img_path, max_width_in=PAGE_TEXT_WIDTH_IN, center=True):
    """Insere imagem; se muito alta p/ a página, fatia em partes de altura de página."""
    with Image.open(img_path) as im:
        w, h = im.size
    disp_h_in = max_width_in * h / w
    if disp_h_in <= PAGE_AVAIL_HEIGHT_IN:
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img_path, width=Inches(max_width_in))
        return
    parts = int(disp_h_in // PAGE_AVAIL_HEIGHT_IN) + (
        1 if disp_h_in % PAGE_AVAIL_HEIGHT_IN else 0
    )
    slice_h = h // parts
    tmpdir = tempfile.mkdtemp(prefix="docxslice_")
    with Image.open(img_path) as im:
        for k in range(parts):
            top = k * slice_h
            bottom = h if k == parts - 1 else (k + 1) * slice_h
            crop = im.crop((0, top, w, bottom))
            sp = os.path.join(tmpdir, f"slice_{k}.png")
            crop.save(sp)
            p = doc.add_paragraph()
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(sp, width=Inches(max_width_in))


def render_blocks(doc, blocks, base_heading=2, skip_mermaid_image=None):
    """Renderiza blocos. skip_mermaid_image: caminho p/ imagem que substitui ```mermaid."""
    for kind, payload in blocks:
        if kind == "heading":
            level, text = payload
            # H2 do arquivo -> Heading base; níveis maiores empurram +1
            doc.add_heading(strip_inline_tags(text), level=min(base_heading + (level - 2), 4))
        elif kind == "para":
            p = doc.add_paragraph()
            add_inline(p, payload)
        elif kind == "table":
            render_table(doc, payload)
        elif kind == "code":
            lang, body = payload
            if lang.lower() == "mermaid" and skip_mermaid_image:
                add_image_fit(doc, skip_mermaid_image, max_width_in=6.5)
            else:
                render_code(doc, lang, body)
        elif kind == "quote":
            render_quote(doc, payload)
        elif kind == "list":
            render_list(doc, payload)

# ================================================================ feature-set specifics

def read_fs_meta(readme_path):
    title, fs_id, dominio = None, None, None
    features = []  # (name, id, filename)
    with open(readme_path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    for ln in lines:
        m = re.match(r"^#\s+Feature Set:\s*(.+)$", ln)
        if m and title is None:
            title = m.group(1).strip()
        # subtítulo do N2: "Major Feature Set: X - `SIGLA-SFS`" (engine 3.0.0) ou "Domínio: X - …" (legado)
        m = re.search(r"(?:Major Feature Set|Domínio):\s*([^-`]+?)\s*-\s*`([A-Z]{3}-[A-Z]{3})`", ln)
        if m and fs_id is None:
            dominio = m.group(1).strip()
            fs_id = m.group(2).strip()
    # tabela de Features: o formato atual, `| [**Nome**](f-x.md) <small>ID</small> | … |`,
    # e o legado, `| **Nome** <small>ID</small> | [f-x.md] … |`
    for ln in lines:
        m = re.match(r"^\|\s*\[\*\*(.+?)\*\*\]\(([^)\s]+\.md)\)\s*<small>([A-Z]{3}-[A-Z]{3}-\d{2})</small>", ln)
        if m:
            features.append((m.group(1).strip(), m.group(3).strip(), m.group(2).strip()))
            continue
        m = re.match(r"^\|\s*\*\*(.+?)\*\*\s*<small>([A-Z]{3}-[A-Z]{3}-\d{2})</small>\s*\|\s*\[([^\]]+)\]", ln)
        if m:
            features.append((m.group(1).strip(), m.group(2).strip(), m.group(3).strip()))
    return title, fs_id, dominio, features


def sistema_da_instancia(fs_dir):
    """"SIGLA — Nome" do global/MASTER.md da instância (o --fs-dir é
    modules/<dominio>/<feature-set>). Campo ainda com o texto do template ([…]) não
    conta; sem MASTER ou sem os dois campos, None."""
    master = os.path.normpath(os.path.join(fs_dir, "..", "..", "..", "global", "MASTER.md"))
    if not os.path.exists(master):
        return None
    with open(master, encoding="utf-8") as fh:
        texto = fh.read()
    def campo(nome):
        m = re.search(rf"^\s*-\s*\*\*{nome}\*\*:\s*`?([^`\n]+?)`?\s*$", texto, re.M)
        valor = m.group(1).strip() if m else ""
        return "" if not valor or valor.startswith("[") else valor
    sigla, nome = campo("Sigla"), campo("Nome")
    if sigla and nome:
        return f"{sigla} — {nome}"
    return sigla or nome or None


def readme_overview_lines(readme_path):
    """Linhas do README até (excl.) o Changelog — a visão geral do N2."""
    out = []
    with open(readme_path, encoding="utf-8") as fh:
        for line in fh:
            if re.match(r"^##\s+Changelog", line):
                break
            out.append(line.rstrip("\n"))
    return out

# ================================================================ main

def style_base(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)


def add_footer_pagenum(doc, label):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(label + "    ·    página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    # campo PAGE
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    r = p.add_run(); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    r._r.append(fldChar1); r._r.append(instr); r._r.append(fldChar2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fs-dir", required=True)
    ap.add_argument("--img-dir", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    readme = os.path.join(args.fs_dir, "README.md")
    title, fs_id, dominio, features = read_fs_meta(readme)
    today = datetime.date.today().isoformat()

    doc = Document()
    style_base(doc)

    # ---- Capa
    sistema = sistema_da_instancia(args.fs_dir)
    if sistema:
        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run(sistema)
        r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x00, 0x55, 0x8C); r.bold = True

    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(f"Especificação de Feature Set (N2)\n{title}")
    r.font.size = Pt(24); r.bold = True
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(f"{fs_id}  ·  Domínio {dominio}")
    r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x55, 0x5F, 0x66)

    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"{len(features)} features · Gerado em {today}\n"
        "Conteúdo negocial (N2 + N3) e protótipos existentes — prova de conceito"
    )
    r.font.size = Pt(9.5); r.italic = True; r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    add_footer_pagenum(doc, f"{fs_id} — {title}")
    doc.add_page_break()

    # ---- Visão geral do N2
    doc.add_heading("Visão geral do Feature Set", level=1)
    overview = readme_overview_lines(readme)
    # remove H1 do README e o "> Nível 2 ..." inicial
    ov_blocks = parse_blocks(overview)
    cleaned = []
    seen_first_heading = False
    for kind, payload in ov_blocks:
        if kind == "heading" and payload[0] == 1 and not seen_first_heading:
            seen_first_heading = True
            continue
        if kind == "quote" and not cleaned and any("Nível 2" in l for l in payload):
            continue
        cleaned.append((kind, payload))
    overview_img = os.path.join(args.img_dir, "overview.png")
    render_blocks(
        doc, cleaned, base_heading=2,
        skip_mermaid_image=overview_img if os.path.exists(overview_img) else None,
    )

    # ---- Features (N3)
    doc.add_page_break()
    doc.add_heading("Especificação das Features (N3)", level=1)
    intro = doc.add_paragraph()
    add_inline(
        intro,
        "As seções a seguir trazem o conteúdo negocial de cada feature — descrição, "
        "regras de negócio, cenários, campos, comportamento de tela e, quando já "
        "contadas, as métricas de tamanho. As seções técnicas (dev-only) e o changelog "
        "são omitidos. Os protótipos são incluídos apenas onde já existem.",
    )

    for name, fid, filename in features:
        doc.add_page_break()
        doc.add_heading(f"{fid} — {name}", level=1)
        fpath = os.path.join(args.fs_dir, filename)
        if not os.path.exists(fpath):
            doc.add_paragraph(f"[arquivo não encontrado: {filename}]")
            continue
        lines = secoes_exportadas(business_lines(fpath))
        blocks = parse_blocks(lines)
        # remove H1 do arquivo + quote "Nível 3"
        cleaned = []
        seen_first_heading = False
        for kind, payload in blocks:
            if kind == "heading" and payload[0] == 1 and not seen_first_heading:
                seen_first_heading = True
                continue
            if kind == "quote" and not cleaned and any("Nível 3" in l for l in payload):
                continue
            cleaned.append((kind, payload))
        render_blocks(doc, cleaned, base_heading=2)

        # protótipos da feature
        protos = sorted(glob.glob(os.path.join(args.img_dir, f"proto-{fid}*.png")))
        if protos:
            doc.add_heading("Protótipo", level=2)
            for k, pp in enumerate(protos):
                parte = f" (tela {k + 1})" if len(protos) > 1 else ""
                cap = doc.add_paragraph()
                rc = cap.add_run(f"Protótipo de tela — {name}{parte}")
                rc.italic = True; rc.font.size = Pt(9); rc.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                add_image_fit(doc, pp)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    doc.save(args.out)
    print(f"OK -> {args.out}  ({len(features)} features)")


if __name__ == "__main__":
    main()
